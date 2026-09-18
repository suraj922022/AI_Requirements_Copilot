import streamlit as st
from google import genai
from dotenv import load_dotenv
import os
from docx import Document
from io import BytesIO
from datetime import datetime

# ==========================
# Load Environment Variables
# ==========================
load_dotenv()

API_KEY = os.getenv("GEMINI_API_KEY")

client = genai.Client(api_key=API_KEY)

# ==========================
# Page Configuration
# ==========================
st.set_page_config(
    page_title="AI Requirements Copilot",
    page_icon="🤖",
    layout="wide"
)

# ==========================
# Session State
# ==========================
if "generated_output" not in st.session_state:
    st.session_state.generated_output = ""

# ==========================
# Header
# ==========================
st.title("🤖 AI Requirements Copilot")

st.markdown("""
Generate and Review Automotive Requirements using Gemini AI.
""")

# ==========================
# Mode Selection
# ==========================
mode = st.radio(
    "Select Mode",
    [
        "Generate Requirements",
        "Review Requirements"
    ]
)

# ==========================
# Domain Selection
# ==========================
domain = st.selectbox(
    "Select Domain",
    [
        "Generic",
        "Bluetooth",
        "Wi-Fi",
        "Android Auto",
        "Apple CarPlay",
        "Connectivity Manager",
        "OTA Update",
        "Navigation",
        "Telematics",
        "Instrument Cluster"
    ]
)

# ======================================================
# GENERATE REQUIREMENTS MODE
# ======================================================
if mode == "Generate Requirements":

    feature_description = st.text_area(
        "Enter Feature Description",
        height=250,
        placeholder="""
Example:

Wireless Android Auto Connection

The infotainment system shall allow users to connect Android Auto wirelessly.

The system shall reconnect automatically after vehicle restart.

The system shall notify the user if connection fails.
"""
    )

    if st.button("Generate Requirements", type="primary"):

        if not feature_description.strip():
            st.warning("Please enter a feature description.")
            st.stop()

        prompt = f"""
You are a Senior Automotive Requirements Engineer.

Domain: {domain}

Generate professional automotive requirements.

IMPORTANT RULES:

- Use automotive terminology.
- Use SHALL statements.
- Do NOT generate markdown tables.
- Do NOT generate HTML tags.
- Do NOT use <br>.
- Do NOT use pipe symbols |.
- Do NOT generate hyperlinks.
- Do NOT generate localhost links.
- Keep requirements atomic and testable.

Generate the following sections.

==================================================
1. FUNCTIONAL REQUIREMENTS
==================================================

FR-001: Requirement

FR-002: Requirement

==================================================
2. NON FUNCTIONAL REQUIREMENTS
==================================================

NFR-001: Requirement

NFR-002: Requirement

Cover:
- Performance
- Reliability
- Security
- Usability

==================================================
3. ACCEPTANCE CRITERIA
==================================================

AC-001: Acceptance Criteria

AC-002: Acceptance Criteria

==================================================
4. TEST CASES
==================================================

TC-001

Description:
...

Preconditions:
...

Steps:
1.
2.
3.

Expected Result:
...

TC-002

Description:
...

Preconditions:
...

Steps:
1.
2.
3.

Expected Result:
...

Feature Description:

{feature_description}
"""

        try:

            with st.spinner("Generating Requirements..."):

                response = client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=prompt
                )

                st.session_state.generated_output = response.text

            st.success("Requirements Generated Successfully")

        except Exception as e:
            st.error(f"Error: {str(e)}")

# ======================================================
# REVIEW REQUIREMENTS MODE
# ======================================================
else:

    requirements_text = st.text_area(
        "Paste Existing Requirements",
        height=300,
        placeholder="""
Example:

System should connect to smartphone automatically.

System should reconnect after restart.
"""
    )

    if st.button("Review Requirements", type="primary"):

        if not requirements_text.strip():
            st.warning("Please enter requirements for review.")
            st.stop()

        prompt = f"""
You are a Senior Automotive Requirements Engineer.

Domain: {domain}

Review the requirements.

Check:

1. Ambiguous wording
2. Missing SHALL statements
3. Testability issues
4. Missing acceptance criteria
5. Duplicate requirements
6. ASPICE requirement quality issues
7. Requirement completeness

For every issue provide:

Issue:
Reason:
Recommendation:
Improved Requirement:

Requirements:

{requirements_text}
"""

        try:

            with st.spinner("Reviewing Requirements..."):

                response = client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=prompt
                )

                st.session_state.generated_output = response.text

            st.success("Review Completed")

        except Exception as e:
            st.error(f"Error: {str(e)}")

# ======================================================
# DISPLAY OUTPUT
# ======================================================
if st.session_state.generated_output:

    st.markdown("---")

    st.subheader("Output")

    st.markdown(st.session_state.generated_output)

    # ==========================
    # Word Export
    # ==========================
    doc = Document()

    doc.add_heading(
        "AI Requirements Copilot Output",
        level=1
    )

    doc.add_paragraph(
        st.session_state.generated_output
    )

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    filename = (
        f"AI_Requirements_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    )

    st.download_button(
        label="📄 Download Word Document",
        data=buffer,
        file_name=filename,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )